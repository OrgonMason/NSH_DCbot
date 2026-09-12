import discord
from discord.ext import commands, tasks
from discord import app_commands
import json
import os
from keep_alive import keep_alive
from datetime import datetime, timezone, timedelta
from dotenv import load_dotenv
from typing import Dict, List, Optional, Tuple
import asyncio
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from collections import defaultdict
import random

load_dotenv()

# ==================== 設定區 ====================
TOKEN = os.getenv("DISCORD_TOKEN")

try:
    from zoneinfo import ZoneInfo
    TZ = ZoneInfo("Asia/Taipei")
except Exception:
    TZ = timezone(timedelta(hours=8))

JOBS = ["鐵衣", "血河", "九靈", "素問", "神相", "龍吟", "碎夢", "玄機", "潮光"]

JOB_EMOJI = {
    "鐵衣": "🟠",
    "血河": "🔴",
    "素問": "🩷",
    "碎夢": "🔵",
    "龍吟": "🟢",
    "玄機": "🟡",
    "九靈": "🟣",
    "神相": "🔷",
    "潮光": "💙",
}

MAX_PLAYERS = 60
MAX_PER_JOB = 14
NUM_TEAMS = 10
SLOTS_PER_TEAM = 6

TEAM_NAMES = [
    "一團一隊", "一團保/拆", "一團保鑣",
    "二團一隊", "二團保/拆", "二團保鑣",
    "三團一隊", "三團二隊", "三團三隊", "三團四隊"
]

REQUIRED_ROLE_NAME = "管理員"

DATA_FILE = "event_data.json"
WEEKDAY_ZH = ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"]
# ================================================

intents = discord.Intents.default()
intents.message_content = True
intents.members = True
intents.guilds = True

bot = commands.Bot(command_prefix="!", intents=intents)


def load_data() -> dict:
    if os.path.exists(DATA_FILE):
        with open(DATA_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_data(data: dict):
    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def get_guild_data(guild_id: int) -> dict:
    data = load_data()
    gid = str(guild_id)
    if gid not in data:
        data[gid] = {
            "event_name": None,
            "event_date": None,
            "opponent": None,
            "deadline": None,
            "deadline_iso": None,
            "channel_id": None,
            "thread_id": None,
            "message_id": None,
            "signups": {},       # 正式報名
            "waitlist": {},      # 候補 {key: info}
            "teams": [[None] * SLOTS_PER_TEAM for _ in range(NUM_TEAMS)],
            "registration_closed": False
        }
        save_data(data)
    g = data[gid]
    if "teams" not in g:
        g["teams"] = [[None] * SLOTS_PER_TEAM for _ in range(NUM_TEAMS)]
    if "waitlist" not in g:
        g["waitlist"] = {}
    return g


def update_guild_data(guild_id: int, guild_data: dict):
    data = load_data()
    data[str(guild_id)] = guild_data
    save_data(data)


def format_datetime_display(dt: datetime) -> str:
    weekday = WEEKDAY_ZH[dt.weekday()]
    return f"{dt.strftime('%Y-%m-%d')}【{weekday}】 {dt.strftime('%H:%M')}"


def parse_yyyymmdd(date_str: str, hour: int, minute: int) -> Optional[datetime]:
    try:
        date_str = date_str.strip().replace("-", "")
        if len(date_str) != 8:
            return None
        year = int(date_str[0:4])
        month = int(date_str[4:6])
        day = int(date_str[6:8])
        return datetime(year, month, day, hour, minute, tzinfo=TZ)
    except Exception:
        return None


def is_deadline_passed(guild_data: dict) -> bool:
    if guild_data.get("registration_closed"):
        return True
    deadline_iso = guild_data.get("deadline_iso")
    if not deadline_iso:
        return False
    try:
        deadline = datetime.fromisoformat(deadline_iso)
        return datetime.now(TZ) >= deadline
    except Exception:
        return False


def get_job_counts(signups: dict) -> Dict[str, int]:
    counts = {job: 0 for job in JOBS}
    for info in signups.values():
        job = info.get("job")
        if job in counts:
            counts[job] += 1
    return counts


def get_waitlist_by_job(waitlist: dict, job: str) -> List[Tuple[str, dict]]:
    """回傳該職業候補名單，依加入順序排序"""
    items = [(k, v) for k, v in waitlist.items() if v.get("job") == job]
    items.sort(key=lambda x: x[1].get("joined_at", ""))
    return items


def count_user_entries(signups: dict, waitlist: dict, user_id: str) -> Tuple[int, int]:
    """回傳 (一般正式+候補次數, 代報正式+候補次數)"""
    normal = 0
    proxy = 0
    for src in (signups, waitlist):
        for key, info in src.items():
            owner = info.get("owner_id", key)
            if owner == user_id:
                if info.get("is_proxy"):
                    proxy += 1
                else:
                    normal += 1
    return normal, proxy


def find_best_slot(guild_data: dict, job: str) -> Optional[Tuple[int, int]]:
    """
    尋找最適合的隊伍位置。
    第一階段：盡量職業不重複（素問最多 2）。
    第二階段：若第一階段找不到（例如鐵衣已超過 10 人），
              允許同隊重複職業，隨機分到仍有空位的隊伍。
    """
    teams = guild_data["teams"]

    def team_job_count(t_idx: int, j: str) -> int:
        return sum(1 for p in teams[t_idx] if p and p.get("job") == j)

    def empty_slots(t_idx: int) -> List[int]:
        return [i for i, p in enumerate(teams[t_idx]) if p is None]

    def team_size(t_idx: int) -> int:
        return sum(1 for p in teams[t_idx] if p is not None)

    # ---------- 第一階段：職業不重複優先 ----------
    candidates = []
    for t in range(NUM_TEAMS):
        empties = empty_slots(t)
        if not empties:
            continue
        if job != "素問" and team_job_count(t, job) >= 1:
            continue
        if job == "素問" and team_job_count(t, "素問") >= 2:
            continue

        if job in ("鐵衣", "血河"):
            preferred = [0] if 0 in empties else empties
        elif job == "素問":
            preferred = [s for s in (4, 5) if s in empties] or empties
        else:
            preferred = empties
        if not preferred:
            continue
        slot = preferred[0]
        score = len(empties) * 10
        if t < 4:
            if job == "鐵衣" and team_job_count(t, "鐵衣") == 0:
                score += 50
            if job == "素問" and team_job_count(t, "素問") < 2:
                score += 40
        elif t in (4, 5):
            if job == "素問" and team_job_count(t, "素問") == 0:
                score += 45
        else:
            if job == "血河" and team_job_count(t, "血河") == 0:
                score += 50
            if job == "素問" and team_job_count(t, "素問") == 0:
                score += 40
        candidates.append((score, t, slot))

    if candidates:
        candidates.sort(reverse=True)
        return candidates[0][1], candidates[0][2]

    # ---------- 第二階段：允許重複職業，分到仍有空位的隊伍 ----------
    fallback = []
    for t in range(NUM_TEAMS):
        empties = empty_slots(t)
        if not empties:
            continue
        # 鐵衣/血河仍優先空著的第 1 位；素問優先 5、6 位
        if job in ("鐵衣", "血河") and 0 in empties:
            slot = 0
        elif job == "素問":
            pref = [s for s in (4, 5) if s in empties]
            slot = pref[0] if pref else random.choice(empties)
        else:
            slot = random.choice(empties)
        # 分數：人數越少的隊伍優先
        score = -team_size(t)
        fallback.append((score, t, slot))

    if not fallback:
        return None

    fallback.sort(reverse=True)
    best_score = fallback[0][0]
    top = [x for x in fallback if x[0] == best_score]
    chosen = random.choice(top)
    return chosen[1], chosen[2]


def place_player(guild_data: dict, key: str, info: dict) -> bool:
    result = find_best_slot(guild_data, info["job"])
    if result is None:
        return False
    t_idx, s_idx = result
    info["team"] = t_idx
    info["slot"] = s_idx
    guild_data["teams"][t_idx][s_idx] = {
        "uid": key,
        "char_name": info["char_name"],
        "job": info["job"],
        "is_proxy": info.get("is_proxy", False),
        "name": info.get("name", "")
    }
    return True


def assign_unassigned_players(guild_data: dict):
    """把尚未分到隊伍的正式報名者補上位置（修正舊資料或先前演算法漏分）"""
    changed = False
    for key, info in list(guild_data.get("signups", {}).items()):
        if info.get("team") is None:
            if place_player(guild_data, key, info):
                changed = True
    return changed


def remove_player_from_teams(guild_data: dict, key: str):
    for t in range(NUM_TEAMS):
        for s in range(SLOTS_PER_TEAM):
            p = guild_data["teams"][t][s]
            if p and p.get("uid") == key:
                guild_data["teams"][t][s] = None


async def promote_from_waitlist(guild: discord.Guild, guild_data: dict, job: str):
    """當正式名額有空缺時，把該職業候補第1位升上正式"""
    waitlist = guild_data.get("waitlist", {})
    items = get_waitlist_by_job(waitlist, job)
    if not items:
        return

    key, info = items[0]
    # 從候補移除
    del waitlist[key]

    # 放到正式
    success = place_player(guild_data, key, info)
    if not success:
        info["team"] = None
        info["slot"] = None
    guild_data["signups"][key] = info
    update_guild_data(guild.id, guild_data)
    await update_thread_status(guild)

    # 私訊通知升上正式
    try:
        owner_id = int(info.get("owner_id", key) if not str(key).startswith("proxy_") else info.get("owner_id"))
        user = guild.get_member(owner_id) or await bot.fetch_user(owner_id)
        if user:
            team_text = ""
            if info.get("team") is not None:
                team_text = f"\n隊伍：{TEAM_NAMES[info['team']]} 第 {info['slot']+1} 位"
            dm = discord.Embed(
                title="🎉 候補轉正通知",
                description=f"你的角色 **{info['char_name']}【{job}】** 已從候補轉為正式報名！",
                color=discord.Color.green()
            )
            dm.add_field(name="活動", value=guild_data.get("event_name", ""), inline=False)
            if team_text:
                dm.add_field(name="分配", value=team_text.strip(), inline=False)
            await user.send(embed=dm)
    except Exception as e:
        print(f"升上正式私訊失敗: {e}")

    # 通知剩下的候補人順位往前
    remaining = get_waitlist_by_job(guild_data.get("waitlist", {}), job)
    for idx, (wkey, winfo) in enumerate(remaining, 1):
        try:
            owner_id = int(winfo.get("owner_id", 0))
            if not owner_id:
                continue
            user = guild.get_member(owner_id) or await bot.fetch_user(owner_id)
            if user:
                await user.send(
                    f"📢 候補順位更新通知\n"
                    f"你的角色 **{winfo['char_name']}【{job}】** 目前候補順位為第 **{idx}** 位。"
                )
        except Exception:
            pass


def format_status_embed(guild_data: dict) -> discord.Embed:
    signups = guild_data.get("signups", {})
    waitlist = guild_data.get("waitlist", {})
    job_counts = get_job_counts(signups)
    remaining = MAX_PLAYERS - len(signups)
    total = len(signups)
    closed = is_deadline_passed(guild_data)

    color = discord.Color.red() if closed or remaining <= 0 else discord.Color.blue()
    embed = discord.Embed(
        title=f"⚔️ {guild_data.get('event_name', '聯賽活動')} 報名狀態",
        color=color
    )
    embed.add_field(name="🏰 對手幫會", value=guild_data.get("opponent") or "未設定", inline=True)
    embed.add_field(name="⏰ 約戰時間", value=guild_data.get("event_date") or "未設定", inline=True)
    embed.add_field(name="📅 報名截止", value=guild_data.get("deadline") or "未設定", inline=True)

    status_text = "🔴 **報名已截止**" if closed else f"🟢 報名進行中（正式剩餘 **{remaining}** 名額）"
    embed.add_field(name="📊 目前正式人數", value=f"{total} / {MAX_PLAYERS}\n{status_text}", inline=False)

    job_lines = []
    for job in JOBS:
        count = job_counts[job]
        emoji = JOB_EMOJI.get(job, "⚪")
        wl_count = len(get_waitlist_by_job(waitlist, job))
        if count >= MAX_PER_JOB:
            status = f"🔴 已滿（候補 {wl_count} 人）"
        else:
            status = f"剩餘 {MAX_PER_JOB - count}"
            if wl_count:
                status += f"（候補 {wl_count} 人）"
        job_lines.append(f"{emoji} **{job}**：{count}/{MAX_PER_JOB}　{status}")
    embed.add_field(name="🗡️ 職業報名狀況", value="\n".join(job_lines), inline=False)

    # 正式名單
    if signups:
        by_job = {job: [] for job in JOBS}
        for info in signups.values():
            job = info.get("job")
            char_name = info.get("char_name") or "未知"
            prefix = "(代報)" if info.get("is_proxy") else ""
            if job in by_job:
                by_job[job].append(f"{prefix}{char_name}【{job}】")
        list_lines = []
        for job in JOBS:
            if by_job[job]:
                emoji = JOB_EMOJI.get(job, "")
                list_lines.append(f"{emoji} **{job}**\n" + "、".join(by_job[job]))
        full_list = "\n\n".join(list_lines)
        if len(full_list) > 900:
            full_list = full_list[:880] + "\n…（已截斷）"
        embed.add_field(name="👥 正式報名人員", value=full_list, inline=False)
    else:
        embed.add_field(name="👥 正式報名人員", value="目前尚無人報名", inline=False)

    # 候補名單
    if waitlist:
        wl_by_job = {job: [] for job in JOBS}
        for job in JOBS:
            items = get_waitlist_by_job(waitlist, job)
            for idx, (_, info) in enumerate(items, 1):
                prefix = "(代報)" if info.get("is_proxy") else ""
                wl_by_job[job].append(f"{idx}.{prefix}{info.get('char_name')}【{job}】")
        wl_lines = []
        for job in JOBS:
            if wl_by_job[job]:
                emoji = JOB_EMOJI.get(job, "")
                wl_lines.append(f"{emoji} **{job}**\n" + "、".join(wl_by_job[job]))
        wl_text = "\n\n".join(wl_lines)
        if len(wl_text) > 900:
            wl_text = wl_text[:880] + "\n…（已截斷）"
        embed.add_field(name="⏳ 候補名單", value=wl_text, inline=False)

    embed.set_footer(text="⛔ 報名已截止" if closed else "請使用下方按鈕進行報名 / 代報")
    embed.timestamp = datetime.now(TZ)
    return embed


def format_team_status_embed(guild_data: dict) -> discord.Embed:
    teams = guild_data.get("teams", [[None]*6 for _ in range(10)])
    job_counts = get_job_counts(guild_data.get("signups", {}))
    embed = discord.Embed(
        title=f"📋 {guild_data.get('event_name', '聯賽')} 隊伍分配狀態",
        color=discord.Color.gold()
    )
    for t_idx in range(NUM_TEAMS):
        team = teams[t_idx]
        lines = []
        for s_idx, p in enumerate(team):
            pos = s_idx + 1
            if p is None:
                lines.append(f"`{pos}.` （空）")
            else:
                emoji = JOB_EMOJI.get(p["job"], "⚪")
                prefix = "(代報)" if p.get("is_proxy") else ""
                lines.append(f"`{pos}.` {emoji} {prefix}{p['char_name']}【{p['job']}】")
        num_emoji = ["1️⃣","2️⃣","3️⃣","4️⃣","5️⃣","6️⃣","7️⃣","8️⃣","9️⃣","🔟"][t_idx]
        embed.add_field(name=f"{num_emoji} {TEAM_NAMES[t_idx]}", value="\n".join(lines), inline=True)

    stats = [f"{JOB_EMOJI.get(j, '⚪')} {j}:{job_counts[j]}" for j in JOBS]
    embed.add_field(name="📊 職業統計", value="　".join(stats), inline=False)
    embed.set_footer(text="使用 /move_member 可調整成員位置")
    embed.timestamp = datetime.now(TZ)
    return embed


async def update_thread_status(guild: discord.Guild):
    guild_data = get_guild_data(guild.id)
    # 自動補上未分配隊伍的成員
    if assign_unassigned_players(guild_data):
        update_guild_data(guild.id, guild_data)
    thread_id = guild_data.get("thread_id")
    message_id = guild_data.get("message_id")
    if not thread_id or not message_id:
        return
    try:
        thread = guild.get_thread(thread_id) or await bot.fetch_channel(thread_id)
    except Exception:
        return

    closed = is_deadline_passed(guild_data)
    if closed and not guild_data.get("registration_closed"):
        guild_data["registration_closed"] = True
        update_guild_data(guild.id, guild_data)

    embed = format_status_embed(guild_data)
    view = SignupView(disabled=closed)
    try:
        msg = await thread.fetch_message(message_id)
        await msg.edit(embed=embed, view=view)
    except Exception:
        new_msg = await thread.send(embed=embed, view=view)
        guild_data["message_id"] = new_msg.id
        update_guild_data(guild.id, guild_data)


class CharacterNameModal(discord.ui.Modal, title="輸入角色名稱"):
    def __init__(self, job: str, is_proxy: bool, select_message: discord.Message = None):
        super().__init__()
        self.job = job
        self.is_proxy = is_proxy
        self.select_message = select_message
        self.char_name_input = discord.ui.TextInput(
            label="角色名稱" + ("（代報）" if is_proxy else ""),
            placeholder="請輸入遊戲內的角色名稱",
            min_length=1,
            max_length=20,
            required=True
        )
        self.add_item(self.char_name_input)

    async def on_submit(self, interaction: discord.Interaction):
        await complete_signup(
            interaction,
            self.job,
            self.char_name_input.value.strip(),
            self.is_proxy,
            select_message=self.select_message
        )



async def cleanup_select_and_response(interaction: discord.Interaction, select_message: discord.Message = None, delay: float = 5):
    """清除成功/失敗提示與下拉選單訊息（兩者一起消失）"""
    # 1) 立刻拿掉下拉選單元件，避免使用者再點
    if select_message is not None:
        try:
            await select_message.edit(content="✅ 處理完成，訊息即將關閉…", view=None)
        except Exception:
            pass

    await asyncio.sleep(delay)

    # 2) 刪除成功/失敗提示
    try:
        await interaction.delete_original_response()
    except Exception:
        pass

    # 3) 刪除下拉選單那則訊息（ephemeral 有時會失敗，多試幾種方式）
    if select_message is not None:
        try:
            await select_message.delete()
        except Exception:
            try:
                # 備援：透過 HTTP 直接刪
                await interaction.client.http.delete_message(select_message.channel.id, select_message.id)
            except Exception:
                try:
                    await select_message.edit(content="（可手動關閉此訊息）", view=None)
                except Exception:
                    pass


async def complete_signup(interaction: discord.Interaction, job: str, char_name: str, is_proxy: bool, select_message: discord.Message = None):
    guild = interaction.guild
    if not guild:
        await interaction.response.send_message("❌ 只能在伺服器內使用。", ephemeral=True)
        return

    guild_data = get_guild_data(guild.id)
    user_id = str(interaction.user.id)

    if is_deadline_passed(guild_data):
        await interaction.response.send_message("⛔ 報名已截止。", ephemeral=True)
        await cleanup_select_and_response(interaction, select_message, delay=4)
        return

    normal_cnt, proxy_cnt = count_user_entries(
        guild_data["signups"], guild_data.get("waitlist", {}), user_id
    )
    # 一般報名限 1 次；代報無上限
    if not is_proxy:
        if normal_cnt >= 1:
            await interaction.response.send_message("❌ 你已經報名過一次了（含候補）。", ephemeral=True)
            await cleanup_select_and_response(interaction, select_message, delay=4)
            return

    job_counts = get_job_counts(guild_data["signups"])
    is_waitlist = job_counts.get(job, 0) >= MAX_PER_JOB

    # 總正式人數仍受 60 限制；候補不受此限制
    if not is_waitlist and len(guild_data["signups"]) >= MAX_PLAYERS:
        # 正式滿了也進候補
        is_waitlist = True

    info = {
        "job": job,
        "char_name": char_name,
        "name": str(interaction.user),
        "is_proxy": is_proxy,
        "owner_id": user_id,
        "joined_at": datetime.now(TZ).isoformat()
    }

    key = user_id if not is_proxy else f"proxy_{user_id}_{int(datetime.now().timestamp())}"

    if is_waitlist:
        # 進入候補
        guild_data.setdefault("waitlist", {})[key] = info
        update_guild_data(guild.id, guild_data)
        await update_thread_status(guild)

        position = len(get_waitlist_by_job(guild_data["waitlist"], job))
        prefix = "(代報)" if is_proxy else ""
        await interaction.response.send_message(
            f"⏳ **已進入候補！**\n"
            f"角色：{prefix}**{char_name}**\n"
            f"職業：**{job}**\n"
            f"目前候補順位：第 **{position}** 位\n"
            f"（有人取消正式名額時會自動遞補並私訊通知你）",
            ephemeral=True
        )
        try:
            dm = discord.Embed(
                title="⏳ 候補成功通知",
                description=f"你的角色 **{prefix}{char_name}【{job}】** 已進入候補名單",
                color=discord.Color.orange()
            )
            dm.add_field(name="候補順位", value=f"第 {position} 位", inline=True)
            dm.add_field(name="活動", value=guild_data.get("event_name", ""), inline=True)
            await interaction.user.send(embed=dm)
        except discord.Forbidden:
            pass
    else:
        # 正式報名
        success = place_player(guild_data, key, info)
        if not success:
            info["team"] = None
            info["slot"] = None
        guild_data["signups"][key] = info
        update_guild_data(guild.id, guild_data)
        await update_thread_status(guild)

        remaining = MAX_PLAYERS - len(guild_data["signups"])
        team_text = ""
        if info.get("team") is not None:
            team_text = f"\n隊伍：{TEAM_NAMES[info['team']]} 第 {info['slot']+1} 位"
        prefix = "(代報)" if is_proxy else ""
        await interaction.response.send_message(
            f"✅ **{'代報' if is_proxy else '報名'}成功！**\n"
            f"角色：{prefix}**{char_name}**\n職業：**{job}**{team_text}\n剩餘正式名額：{remaining}",
            ephemeral=True
        )
        try:
            thread_id = guild_data.get("thread_id")
            thread_link = f"https://discord.com/channels/{guild.id}/{thread_id}" if thread_id else ""
            dm = discord.Embed(
                title="🎉 報名成功通知" if not is_proxy else "📝 代報成功通知",
                description=f"{'你已成功報名' if not is_proxy else '你已成功代報'} **{guild_data['event_name']}**",
                color=discord.Color.green()
            )
            dm.add_field(name="角色名稱", value=f"{prefix}{char_name}", inline=True)
            dm.add_field(name="職業", value=job, inline=True)
            if team_text:
                dm.add_field(name="分配隊伍", value=team_text.strip(), inline=False)
            if thread_link:
                dm.add_field(name="📢 討論串", value=f"[點此前往]({thread_link})", inline=False)
            await interaction.user.send(embed=dm)
        except discord.Forbidden:
            pass

    await cleanup_select_and_response(interaction, select_message, delay=5)


async def cancel_one_entry(interaction: discord.Interaction, src_name: str, key: str, select_message: discord.Message = None):
    """取消單一筆報名/代報（正式或候補）"""
    guild = interaction.guild
    guild_data = get_guild_data(guild.id)
    src = guild_data.get(src_name, {})
    info = src.get(key)
    if not info:
        await interaction.response.send_message("❌ 找不到該筆資料，可能已被取消。", ephemeral=True)
        await cleanup_select_and_response(interaction, select_message, delay=4)
        return

    char_name = info.get("char_name", "未知")
    job = info.get("job", "")
    is_proxy = info.get("is_proxy", False)
    kind = "代報" if is_proxy else "報名"

    src.pop(key, None)
    if src_name == "signups":
        remove_player_from_teams(guild_data, key)
        update_guild_data(guild.id, guild_data)
        if job:
            await promote_from_waitlist(guild, guild_data, job)
    else:
        update_guild_data(guild.id, guild_data)

    await update_thread_status(guild)

    label = f"{'(代報)' if is_proxy else ''}{char_name}【{job}】"
    status = "正式" if src_name == "signups" else "候補"
    await interaction.response.send_message(
        f"✅ 已取消{kind}：**{label}**（{status}）",
        ephemeral=True
    )

    try:
        dm = discord.Embed(
            title=f"❌ 取消{kind}通知",
            color=discord.Color.orange(),
            description=f"你已取消 **{guild_data.get('event_name')}** 的{kind}\n{label}（{status}）"
        )
        await interaction.user.send(embed=dm)
    except discord.Forbidden:
        pass

    # select_message：取消代報時的下拉選單訊息；若無則用 interaction.message
    msg = select_message or interaction.message
    await cleanup_select_and_response(interaction, msg, delay=5)


async def handle_cancel_normal(interaction: discord.Interaction):
    """取消自己的一般報名（正式或候補，通常只有一筆）"""
    guild = interaction.guild
    if not guild:
        return
    guild_data = get_guild_data(guild.id)
    user_id = str(interaction.user.id)

    if is_deadline_passed(guild_data):
        await interaction.response.send_message("⛔ 報名已截止，無法取消。", ephemeral=True)
        return

    targets = []
    for src_name in ("signups", "waitlist"):
        src = guild_data.get(src_name, {})
        for key, info in list(src.items()):
            owner = info.get("owner_id", key)
            if owner == user_id and not info.get("is_proxy", False):
                targets.append((src_name, key, info))

    if not targets:
        await interaction.response.send_message("❌ 你目前沒有任何報名紀錄。", ephemeral=True)
        return

    # 一般報名通常只有一筆，直接取消第一筆
    src_name, key, _info = targets[0]
    await cancel_one_entry(interaction, src_name, key)


async def handle_cancel_proxy_start(interaction: discord.Interaction):
    """點「取消代報」後，列出可選的代報成員讓使用者指定"""
    guild = interaction.guild
    if not guild:
        return
    guild_data = get_guild_data(guild.id)
    user_id = str(interaction.user.id)

    if is_deadline_passed(guild_data):
        await interaction.response.send_message("⛔ 報名已截止，無法取消。", ephemeral=True)
        return

    options = []
    # value 格式：src_name|key
    for src_name, src_label in (("signups", "正式"), ("waitlist", "候補")):
        src = guild_data.get(src_name, {})
        for key, info in src.items():
            owner = info.get("owner_id", key)
            if owner != user_id or not info.get("is_proxy", False):
                continue
            char_name = info.get("char_name", "未知")
            job = info.get("job", "")
            label = f"{char_name}【{job}】（{src_label}）"
            # Select option label max 100, value max 100
            value = f"{src_name}|{key}"
            if len(value) > 100:
                value = value[:100]
            options.append(discord.SelectOption(
                label=label[:100],
                value=value,
                description=f"{src_label}・點選以取消此代報"
            ))

    if not options:
        await interaction.response.send_message("❌ 你目前沒有任何代報紀錄。", ephemeral=True)
        return

    # Discord 一個 select 最多 25 個選項
    if len(options) > 25:
        options = options[:25]

    view = CancelProxySelectView(options)
    await interaction.response.send_message(
        "請選擇要**取消的代報成員**：",
        view=view,
        ephemeral=True
    )


class CancelProxySelectView(discord.ui.View):
    def __init__(self, options: list):
        super().__init__(timeout=60)
        select = discord.ui.Select(
            placeholder="選擇要取消的代報角色…",
            options=options,
            custom_id="cancel_proxy_select"
        )
        select.callback = self.on_select
        self.add_item(select)

    async def on_select(self, interaction: discord.Interaction):
        value = interaction.data["values"][0]
        if "|" not in value:
            await interaction.response.send_message("❌ 資料格式錯誤", ephemeral=True)
            return
        src_name, key = value.split("|", 1)
        if src_name not in ("signups", "waitlist"):
            await interaction.response.send_message("❌ 資料錯誤", ephemeral=True)
            return
        # interaction.message = 這則「選擇要取消的代報」下拉選單訊息
        await cancel_one_entry(interaction, src_name, key, select_message=interaction.message)


class JobSelectView(discord.ui.View):
    def __init__(self, is_proxy: bool):
        super().__init__(timeout=60)
        self.is_proxy = is_proxy
        options = [discord.SelectOption(label=job, value=job) for job in JOBS]
        select = discord.ui.Select(placeholder="請選擇職業", options=options)
        select.callback = self.on_select
        self.add_item(select)

    async def on_select(self, interaction: discord.Interaction):
        job = interaction.data["values"][0]
        modal = CharacterNameModal(job, self.is_proxy, select_message=interaction.message)
        await interaction.response.send_modal(modal)


class SignupView(discord.ui.View):
    def __init__(self, disabled: bool = False):
        super().__init__(timeout=None)

        btn_signup = discord.ui.Button(label="報名", style=discord.ButtonStyle.primary,
                                       custom_id="btn_signup", emoji="📝", disabled=disabled)
        btn_signup.callback = self.signup_callback
        self.add_item(btn_signup)

        btn_proxy = discord.ui.Button(label="代報", style=discord.ButtonStyle.secondary,
                                      custom_id="btn_proxy", emoji="👥", disabled=disabled)
        btn_proxy.callback = self.proxy_callback
        self.add_item(btn_proxy)

        btn_cancel = discord.ui.Button(label="取消報名", style=discord.ButtonStyle.danger,
                                       custom_id="btn_cancel", emoji="⚠️", disabled=disabled)
        btn_cancel.callback = self.cancel_callback
        self.add_item(btn_cancel)

        btn_cancel_proxy = discord.ui.Button(label="取消代報", style=discord.ButtonStyle.danger,
                                             custom_id="btn_cancel_proxy", emoji="⚠️", disabled=disabled)
        btn_cancel_proxy.callback = self.cancel_proxy_callback
        self.add_item(btn_cancel_proxy)

        btn_refresh = discord.ui.Button(label="重新整理", style=discord.ButtonStyle.secondary,
                                        custom_id="btn_refresh", emoji="🔄")
        btn_refresh.callback = self.refresh_callback
        self.add_item(btn_refresh)

    async def signup_callback(self, interaction: discord.Interaction):
        if is_deadline_passed(get_guild_data(interaction.guild.id)):
            await interaction.response.send_message("⛔ 報名已截止。", ephemeral=True)
            return
        await interaction.response.send_message("請選擇要報名的職業：",
                                                view=JobSelectView(is_proxy=False), ephemeral=True)

    async def proxy_callback(self, interaction: discord.Interaction):
        if is_deadline_passed(get_guild_data(interaction.guild.id)):
            await interaction.response.send_message("⛔ 報名已截止。", ephemeral=True)
            return
        await interaction.response.send_message("請選擇要**代報**的職業：",
                                                view=JobSelectView(is_proxy=True), ephemeral=True)

    async def cancel_callback(self, interaction: discord.Interaction):
        await handle_cancel_normal(interaction)

    async def cancel_proxy_callback(self, interaction: discord.Interaction):
        await handle_cancel_proxy_start(interaction)

    async def refresh_callback(self, interaction: discord.Interaction):
        await update_thread_status(interaction.guild)
        await interaction.response.send_message("✅ 已重新整理", ephemeral=True)
        await asyncio.sleep(3)
        try:
            await interaction.delete_original_response()
        except Exception:
            pass


def has_required_role():
    async def predicate(interaction: discord.Interaction) -> bool:
        if interaction.user.guild_permissions.administrator:
            return True
        role = discord.utils.get(interaction.guild.roles, name=REQUIRED_ROLE_NAME)
        if role and role in interaction.user.roles:
            return True
        await interaction.response.send_message(
            f"❌ 你需要擁有 **{REQUIRED_ROLE_NAME}** 身分組才能使用此指令。", ephemeral=True)
        return False
    return app_commands.check(predicate)


HOUR_CHOICES = [app_commands.Choice(name=f"{h:02d} 時", value=h) for h in range(24)]
MINUTE_CHOICES = [
    app_commands.Choice(name="00 分（整點）", value=0),
    app_commands.Choice(name="30 分", value=30),
]


@bot.tree.command(name="new_event", description="【管理員】建立新的聯賽活動")
@app_commands.describe(
    名稱="活動名稱",
    對手幫會="對手幫會名稱",
    約戰日期="約戰日期（格式 YYYYMMDD）",
    約戰時="約戰小時",
    約戰分="約戰分鐘",
    截止日期="報名截止日期（格式 YYYYMMDD）",
    截止時="截止小時",
    截止分="截止分鐘"
)
@app_commands.choices(約戰時=HOUR_CHOICES, 約戰分=MINUTE_CHOICES, 截止時=HOUR_CHOICES, 截止分=MINUTE_CHOICES)
@app_commands.default_permissions(administrator=True)
@has_required_role()
async def new_event(
    interaction: discord.Interaction,
    名稱: str, 對手幫會: str,
    約戰日期: str, 約戰時: app_commands.Choice[int], 約戰分: app_commands.Choice[int],
    截止日期: str, 截止時: app_commands.Choice[int], 截止分: app_commands.Choice[int]
):
    await interaction.response.defer(ephemeral=True)

    battle_dt = parse_yyyymmdd(約戰日期, 約戰時.value, 約戰分.value)
    if not battle_dt:
        await interaction.followup.send("❌ 約戰日期格式錯誤，請使用 YYYYMMDD（例如 20260920）", ephemeral=True)
        return
    deadline_dt = parse_yyyymmdd(截止日期, 截止時.value, 截止分.value)
    if not deadline_dt:
        await interaction.followup.send("❌ 截止日期格式錯誤，請使用 YYYYMMDD", ephemeral=True)
        return
    if deadline_dt >= battle_dt:
        await interaction.followup.send("❌ 報名截止時間必須早於約戰時間", ephemeral=True)
        return

    battle_display = format_datetime_display(battle_dt)
    deadline_display = format_datetime_display(deadline_dt)

    guild = interaction.guild
    channel = interaction.channel
    if isinstance(channel, discord.Thread):
        channel = channel.parent
    if not isinstance(channel, discord.TextChannel):
        await interaction.followup.send("❌ 請在文字頻道使用", ephemeral=True)
        return

    guild_data = get_guild_data(guild.id)
    guild_data.update({
        "event_name": 名稱,
        "opponent": 對手幫會,
        "event_date": battle_display,
        "deadline": deadline_display,
        "deadline_iso": deadline_dt.isoformat(),
        "signups": {},
        "waitlist": {},
        "teams": [[None] * SLOTS_PER_TEAM for _ in range(NUM_TEAMS)],
        "registration_closed": False,
        "channel_id": channel.id
    })

    try:
        thread = await channel.create_thread(
            name=f"【{名稱}】vs {對手幫會}｜{battle_display}",
            type=discord.ChannelType.public_thread,
            auto_archive_duration=10080
        )
        guild_data["thread_id"] = thread.id

        try:
            await thread.set_permissions(guild.default_role, send_messages=False)
        except Exception as e:
            print(f"設定討論串權限時發生問題（可忽略）: {e}")

        embed = format_status_embed(guild_data)
        view = SignupView(disabled=False)
        msg = await thread.send(
            content=(
                f"📢 **{名稱} 報名討論串**\n"
                f"🏰 對手：{對手幫會}\n"
                f"⏰ 約戰：{battle_display}\n"
                f"📅 截止：{deadline_display}\n\n"
                f"請使用下方按鈕「報名」或「代報」\n"
                f"（本討論串已禁止一般訊息，僅能透過按鈕操作）"
            ),
            embed=embed, view=view
        )
        guild_data["message_id"] = msg.id
        update_guild_data(guild.id, guild_data)
        await interaction.followup.send(f"✅ 活動已建立！\n討論串：{thread.mention}", ephemeral=True)
    except Exception as e:
        await interaction.followup.send(f"❌ 建立失敗：{e}", ephemeral=True)


@bot.tree.command(name="export_excel", description="【管理員】匯出報名名單 Excel（含分組）")
@app_commands.default_permissions(administrator=True)
@has_required_role()
async def export_excel(interaction: discord.Interaction):
    await interaction.response.defer(ephemeral=True)
    guild_data = get_guild_data(interaction.guild.id)
    signups = guild_data.get("signups", {})
    waitlist = guild_data.get("waitlist", {})
    teams = guild_data.get("teams", [])

    if not signups and not waitlist:
        await interaction.followup.send("❌ 目前沒有報名資料", ephemeral=True)
        return

    wb = Workbook()
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    ws = wb.active
    ws.title = "分組名單"
    ws.merge_cells('A1:E1')
    ws['A1'] = f"{guild_data.get('event_name')} vs {guild_data.get('opponent')}｜{guild_data.get('event_date')}"
    ws['A1'].font = Font(bold=True, size=14)

    headers = ["隊伍", "位置", "角色名稱", "職業", "代報"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin

    row = 4
    for t_idx, team in enumerate(teams):
        for s_idx, p in enumerate(team):
            if p is None:
                values = [TEAM_NAMES[t_idx], s_idx+1, "（空）", "", ""]
            else:
                values = [TEAM_NAMES[t_idx], s_idx+1, p.get("char_name", ""), p.get("job", ""),
                          "是" if p.get("is_proxy") else "否"]
            for c, v in enumerate(values, 1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.border = thin
                cell.alignment = Alignment(horizontal="center")
            row += 1
        row += 1

    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 16

    ws2 = wb.create_sheet("完整名單")
    for col, h in enumerate(["序號", "角色名稱", "職業", "代報", "狀態", "Discord名稱", "ID"], 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font

    idx = 1
    for key, info in signups.items():
        ws2.append([idx, info.get("char_name"), info.get("job"),
                    "是" if info.get("is_proxy") else "否", "正式", info.get("name"), key])
        idx += 1
    for key, info in waitlist.items():
        ws2.append([idx, info.get("char_name"), info.get("job"),
                    "是" if info.get("is_proxy") else "否", "候補", info.get("name"), key])
        idx += 1

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"{guild_data.get('event_name', 'event')}_名單_{datetime.now(TZ).strftime('%Y%m%d_%H%M')}.xlsx"
    await interaction.followup.send(
        content=f"📊 已匯出（正式 {len(signups)} + 候補 {len(waitlist)} 人）",
        file=discord.File(fp=buffer, filename=filename), ephemeral=True
    )


@bot.tree.command(name="event_status", description="查看目前隊伍分配狀態")
@has_required_role()
async def event_status(interaction: discord.Interaction):
    guild_data = get_guild_data(interaction.guild.id)
    if not guild_data.get("event_name"):
        await interaction.response.send_message("❌ 目前沒有進行中的活動", ephemeral=True)
        return
    embed = format_team_status_embed(guild_data)
    await interaction.response.send_message(embed=embed, ephemeral=True)


@bot.tree.command(name="close_event", description="【管理員】強制關閉報名")
@app_commands.default_permissions(administrator=True)
@has_required_role()
async def close_event(interaction: discord.Interaction):
    guild_data = get_guild_data(interaction.guild.id)
    guild_data["registration_closed"] = True
    update_guild_data(interaction.guild.id, guild_data)
    await update_thread_status(interaction.guild)
    await interaction.response.send_message("✅ 已強制關閉報名", ephemeral=True)



@bot.tree.command(name="who_proxy", description="查詢某個角色是由誰代報的")
@app_commands.describe(角色名稱="要查詢的角色名稱")
@has_required_role()
async def who_proxy(interaction: discord.Interaction, 角色名稱: str):
    guild_data = get_guild_data(interaction.guild.id)
    signups = guild_data.get("signups", {})
    waitlist = guild_data.get("waitlist", {})

    found = None
    source = None
    for src_name, src in (("正式", signups), ("候補", waitlist)):
        for key, info in src.items():
            if info.get("char_name") == 角色名稱:
                found = info
                source = src_name
                break
        if found:
            break

    if not found:
        await interaction.response.send_message(f"❌ 無該成員「{角色名稱}」", ephemeral=True)
        return

    if not found.get("is_proxy"):
        await interaction.response.send_message(
            f"ℹ️ **{角色名稱}【{found.get('job')}】** 為正式自行報名成員（非代報）。\n狀態：{source}",
            ephemeral=True
        )
        return

    proxy_by = found.get("name", "未知")
    owner_id = found.get("owner_id", "")
    await interaction.response.send_message(
        f"👥 **{角色名稱}【{found.get('job')}】** 是由 **{proxy_by}** 代報的。\n"
        f"狀態：{source}\n"
        f"代報者 Discord ID：`{owner_id}`",
        ephemeral=True
    )


@bot.tree.command(name="move_member", description="【管理員】調整成員到指定隊伍與位置")
@app_commands.describe(
    角色名稱="要移動的角色名稱",
    目標隊伍="1~10（對應一團一隊～三團四隊）",
    目標位置="1~6"
)
@app_commands.default_permissions(administrator=True)
@has_required_role()
async def move_member(
    interaction: discord.Interaction,
    角色名稱: str,
    目標隊伍: app_commands.Range[int, 1, 10],
    目標位置: app_commands.Range[int, 1, 6]
):
    await interaction.response.defer(ephemeral=True)
    guild_data = get_guild_data(interaction.guild.id)
    teams = guild_data["teams"]
    signups = guild_data["signups"]

    found_key = None
    found_info = None
    for key, info in signups.items():
        if info.get("char_name") == 角色名稱:
            found_key = key
            found_info = info
            break

    if not found_info:
        await interaction.followup.send(f"❌ 找不到角色「{角色名稱}」（僅能移動正式名單）", ephemeral=True)
        return

    t_idx = 目標隊伍 - 1
    s_idx = 目標位置 - 1

    old_p = teams[t_idx][s_idx]
    if old_p:
        for k, inf in signups.items():
            if inf.get("char_name") == old_p.get("char_name"):
                inf["team"] = None
                inf["slot"] = None

    remove_player_from_teams(guild_data, found_key)
    teams[t_idx][s_idx] = {
        "uid": found_key,
        "char_name": found_info["char_name"],
        "job": found_info["job"],
        "is_proxy": found_info.get("is_proxy", False),
        "name": found_info.get("name", "")
    }
    found_info["team"] = t_idx
    found_info["slot"] = s_idx

    update_guild_data(interaction.guild.id, guild_data)
    await update_thread_status(interaction.guild)

    await interaction.followup.send(
        f"✅ 已將 **{角色名稱}** 移動到 **{TEAM_NAMES[t_idx]}** 的第 **{目標位置}** 位置",
        ephemeral=True
    )


@tasks.loop(minutes=1)
async def check_deadline_task():
    data = load_data()
    for gid, gdata in data.items():
        if gdata.get("registration_closed"):
            continue
        if is_deadline_passed(gdata):
            gdata["registration_closed"] = True
            update_guild_data(int(gid), gdata)
            guild = bot.get_guild(int(gid))
            if guild:
                await update_thread_status(guild)


@bot.event
async def on_ready():
    print(f"✅ 已登入為 {bot.user}")
    bot.add_view(SignupView(disabled=False))
    try:
        synced = await bot.tree.sync()
        print(f"✅ 同步 {len(synced)} 個指令：")
        for cmd in synced:
            print(f"   /{cmd.name}")
    except Exception as e:
        print(f"❌ 同步指令失敗: {e}")
    if not check_deadline_task.is_running():
        check_deadline_task.start()


if __name__ == "__main__":
    keep_alive()
    print("✅ keep_alive 已啟動")

    token = os.getenv("DISCORD_TOKEN") or ""
    if not token:
        raise Exception("Please set DISCORD_TOKEN")

    try:
        bot.run(token)
    except discord.HTTPException as e:
        if e.status == 429:
            print("The Discord servers denied the connection for making too many requests")
            print("Get help from https://stackoverflow.com/questions/66724687/in-discord-py-how-to-solve-the-error-for-toomanyrequests")
        else:
            raise e
